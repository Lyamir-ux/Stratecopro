#!/usr/bin/env python3
"""Lecture d'un classeur de plan de financement (PF) avant conversion.

Affiche, pour chaque onglet : son état (visible / masqué), ses colonnes et
lignes masquées, ses cellules fusionnées, puis chaque ligne non vide avec,
pour chaque cellule, la valeur calculée et la formule (« =D21+D22 »).

Les formules disent ce que les valeurs taisent : lignes retenues dans
l'assiette MaPrimeRénov' (« Total travaux HT retenu = D21+D22+D24 »), TVA
(« =4410*1.2 »), pourcentages (« =D35*3/100 »), coefficients des aides
(« =D36*0.45*0.9 »).

Usage :
    python lire_classeur.py CLASSEUR.xlsx                  # tous les onglets
    python lire_classeur.py CLASSEUR.xlsx --onglets        # liste des onglets seulement
    python lire_classeur.py CLASSEUR.xlsx --onglet "PF définitif"   # un onglet (nom partiel accepté)
    python lire_classeur.py CLASSEUR.xlsx --max-lignes 400 --sans-masques
"""
from __future__ import annotations

import argparse
import datetime as dt
import sys

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl est requis : pip install openpyxl")


def fmt_valeur(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VRAI" if v else "FAUX"
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return repr(round(v, 6))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    s = str(v).replace("\n", " / ").strip()
    return s


def colonnes_masquees(ws) -> list[str]:
    out = []
    for lettre, dim in ws.column_dimensions.items():
        if dim.hidden:
            # une dimension peut couvrir plusieurs colonnes (min..max)
            if dim.min and dim.max and dim.max > dim.min:
                out.append(f"{get_column_letter(dim.min)}:{get_column_letter(dim.max)}")
            else:
                out.append(lettre)
    return out


def lignes_masquees(ws) -> list[int]:
    return sorted(r for r, dim in ws.row_dimensions.items() if dim.hidden)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("classeur")
    p.add_argument("--onglets", action="store_true", help="liste des onglets seulement")
    p.add_argument("--onglet", help="nom (ou partie du nom) de l'onglet à afficher")
    p.add_argument("--max-lignes", type=int, default=300, help="lignes affichées par onglet (défaut 300)")
    p.add_argument("--max-colonnes", type=int, default=20, help="colonnes lues par ligne (défaut 20, A à T)")
    p.add_argument("--sans-masques", action="store_true", help="n'affiche pas les onglets masqués")
    a = p.parse_args()

    # Deux lectures : formules (data_only=False) et dernières valeurs calculées
    # par Excel (data_only=True). Un classeur jamais ouvert dans Excel après
    # génération par un script n'a pas de valeurs calculées : elles s'affichent vides.
    wb_f = openpyxl.load_workbook(a.classeur, data_only=False)
    wb_v = openpyxl.load_workbook(a.classeur, data_only=True)

    print(f"Classeur : {a.classeur}")
    print("Onglets :")
    for ws in wb_f.worksheets:
        etat = {"visible": "visible", "hidden": "MASQUÉ", "veryHidden": "TRÈS MASQUÉ"}.get(ws.sheet_state, ws.sheet_state)
        print(f"  - « {ws.title} » ({etat}, {ws.max_row} lignes x {ws.max_column} colonnes)")
    if a.onglets:
        return

    for ws_f in wb_f.worksheets:
        if a.onglet and a.onglet.lower() not in ws_f.title.lower():
            continue
        if a.sans_masques and ws_f.sheet_state != "visible":
            continue
        ws_v = wb_v[ws_f.title]
        print()
        print("=" * 100)
        etat = "visible" if ws_f.sheet_state == "visible" else f"MASQUÉ ({ws_f.sheet_state})"
        print(f"ONGLET « {ws_f.title} » - {etat}")
        cm = colonnes_masquees(ws_f)
        if cm:
            print(f"Colonnes masquées : {', '.join(cm)}")
        lm = lignes_masquees(ws_f)
        if lm:
            apercu = ", ".join(str(r) for r in lm[:40]) + (" ..." if len(lm) > 40 else "")
            print(f"Lignes masquées : {apercu}")
        fusions = [str(m) for m in ws_f.merged_cells.ranges]
        if fusions:
            print(f"Cellules fusionnées : {', '.join(fusions[:80])}" + (" ..." if len(fusions) > 80 else ""))
        print("-" * 100)
        n = 0
        max_col = min(ws_f.max_column, a.max_colonnes)
        for r in range(1, ws_f.max_row + 1):
            morceaux = []
            for c in range(1, max_col + 1):
                cf = ws_f.cell(r, c)
                v = ws_v.cell(r, c).value
                f = cf.value if isinstance(cf.value, str) and cf.value.startswith("=") else None
                if f is None and cf.data_type == "f":
                    f = str(cf.value)
                if v is None and f is None:
                    if cf.value is not None and not isinstance(cf.value, str):
                        v = cf.value
                    elif isinstance(cf.value, str) and not cf.value.startswith("="):
                        v = cf.value
                if v is None and f is None:
                    continue
                txt = fmt_valeur(v)
                if f:
                    txt = f"{txt} [{f}]" if txt != "" else f"(non calculé) [{f}]"
                morceaux.append(f"{get_column_letter(c)}: {txt}")
            if not morceaux:
                continue
            masque = " (masquée)" if r in lm else ""
            print(f"L{r}{masque} | " + " | ".join(morceaux))
            n += 1
            if n >= a.max_lignes:
                print(f"... (affichage limité à {a.max_lignes} lignes, utiliser --max-lignes)")
                break


if __name__ == "__main__":
    main()
